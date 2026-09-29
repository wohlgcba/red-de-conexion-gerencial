"""Provision RCG Auth users from the private XLSX. Dry-run unless --execute.

The workbook is local and gitignored. CUIT values are transmitted only to
Supabase Auth as initial passwords; they are never printed or stored in SQL.
Existing Auth accounts (including Hub users) are never changed.
"""

from __future__ import annotations

import argparse
import json
import shutil
import subprocess
import time
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.parse import urlencode
from urllib.request import Request, urlopen

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
PROJECT_REF = "pnhskmlejdaklwkvasxp"  # BASE GCBA only
BASE_URL = f"https://{PROJECT_REF}.supabase.co/auth/v1/admin/users"
SHEET_NAME = "Respuestas de formulario 1"
SQL_LINK = ROOT / "supabase" / "rcg_link_new_accounts.sql"


def valid_cuit(value: object) -> str | None:
    if not isinstance(value, (int, float)) or int(value) != value:
        return None
    digits = str(int(value))
    if len(digits) != 11:
        return None
    factors = (5, 4, 3, 2, 7, 6, 5, 4, 3, 2)
    expected = (11 - sum(int(digit) * factor for digit, factor in zip(digits[:10], factors)) % 11) % 11
    return digits if expected == int(digits[-1]) else None


def load_people() -> tuple[list[tuple[int, str, str]], list[int]]:
    workbooks = list(ROOT.glob("*.xlsx"))
    if len(workbooks) != 1:
        raise RuntimeError("Se esperaba exactamente un XLSX local en la raíz.")
    workbook = load_workbook(workbooks[0], read_only=True, data_only=True)
    try:
        sheet = workbook[SHEET_NAME]
        headers = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        if "Email" not in str(headers[9]) or "Cuit" not in str(headers[11]):
            raise RuntimeError("Las columnas del XLSX no coinciden con las esperadas.")
        valid: list[tuple[int, str, str]] = []
        invalid: list[int] = []
        seen: set[str] = set()
        for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
            if not any(cell is not None and str(cell).strip() for cell in row):
                continue
            email = str(row[9] or "").strip().lower()
            if not email or email.count("@") != 1 or email in seen:
                raise RuntimeError(f"Correo inválido o duplicado en fila {row_number}.")
            seen.add(email)
            cuit = valid_cuit(row[11])
            if cuit:
                valid.append((row_number, email, cuit))
            else:
                invalid.append(row_number)
        return valid, invalid
    finally:
        workbook.close()


def get_service_key(cli: str) -> str:
    result = subprocess.run(
        [cli, "--yes", "supabase@latest", "projects", "api-keys", "--project-ref", PROJECT_REF, "--output", "json", "--agent", "no"],
        cwd=ROOT, capture_output=True, text=True, check=False,
    )
    if result.returncode:
        raise RuntimeError("La CLI no tiene acceso a BASE GCBA. Iniciá sesión con la cuenta correcta.")
    keys = json.loads(result.stdout)
    key = next((entry.get("api_key") for entry in keys if entry.get("name") == "service_role"), None)
    if not key or len(key) < 40:
        raise RuntimeError("No se encontró la clave administrativa de BASE GCBA.")
    return key


def api_request(key: str, method: str, url: str, body: dict | None = None) -> dict:
    payload = json.dumps(body).encode("utf-8") if body is not None else None
    request = Request(url, data=payload, method=method, headers={
        "apikey": key, "Authorization": f"Bearer {key}", "Content-Type": "application/json",
    })
    for attempt in range(4):
        try:
            with urlopen(request, timeout=30) as response:
                return json.load(response)
        except HTTPError as error:
            if error.code in (429, 500, 502, 503, 504) and attempt < 3:
                time.sleep(2 ** attempt)
                continue
            raise RuntimeError(f"Falló la API administrativa de Supabase (HTTP {error.code}).") from None
        except URLError:
            if attempt < 3:
                time.sleep(2 ** attempt)
                continue
            raise RuntimeError("No se pudo conectar con Supabase Auth.") from None
    raise RuntimeError("No se pudo completar la solicitud a Supabase Auth.")


def existing_emails(key: str) -> set[str]:
    result: set[str] = set()
    page = 1
    while True:
        query = urlencode({"page": page, "per_page": 1000})
        users = api_request(key, "GET", f"{BASE_URL}?{query}").get("users", [])
        result.update(str(user.get("email") or "").strip().lower() for user in users)
        if len(users) < 1000:
            return result
        page += 1


def link_new_accounts(cli: str) -> None:
    completed = subprocess.run(
        [cli, "--yes", "supabase@latest", "db", "query", "--linked", "--project-ref", PROJECT_REF,
         "--file", str(SQL_LINK), "--agent", "no", "--output-format", "text"],
        cwd=ROOT, capture_output=True, text=True, check=False,
    )
    if completed.returncode:
        raise RuntimeError("Las cuentas se crearon, pero falló el vínculo con el directorio. Repetí el script para completarlo.")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--execute", action="store_true", help="Crear cuentas nuevas (sin este flag solo informa conteos).")
    parser.add_argument("--limit", type=int, default=0, help="Límite opcional para una prueba controlada.")
    args = parser.parse_args()
    valid, invalid = load_people()
    print(f"Filas con CUIT válido: {len(valid)}. Filas pendientes de corrección: {invalid}.")
    if not args.execute:
        print("Simulación local: no se crearon cuentas ni se modificó la base.")
        return
    cli = shutil.which("npx")
    if not cli:
        raise RuntimeError("No se encontró npx.")
    key = get_service_key(cli)
    existing = existing_emails(key)
    new = [item for item in valid if item[1] not in existing]
    if args.limit:
        new = new[:args.limit]
    print(f"Cuentas existentes que se conservan: {len(valid) - len([item for item in valid if item[1] not in existing])}. Cuentas nuevas a crear: {len(new)}.")
    created = 0
    try:
        for row_number, email, cuit in new:
            try:
                api_request(key, "POST", BASE_URL, {
                    "email": email, "password": cuit, "email_confirm": True,
                    "app_metadata": {"rcg_initial_cuit": True},
                })
            except RuntimeError as error:
                raise RuntimeError(f"Se detuvo en la fila {row_number}: {error}") from None
            created += 1
            if created % 25 == 0:
                link_new_accounts(cli)
                print(f"Cuentas creadas y vinculadas: {created}.")
    finally:
        link_new_accounts(cli)
    print(f"Finalizado. Cuentas nuevas creadas: {created}; las existentes no se modificaron.")


if __name__ == "__main__":
    try:
        main()
    except RuntimeError as error:
        raise SystemExit(str(error)) from None
