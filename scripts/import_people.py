"""Import the directory worksheet into the private BASE GCBA schema.

Only the active directory fields are copied. The workbook, generated SQL and
any credentials must never be committed to the public repository.
Requires: Python with openpyxl, Node/npm and an authenticated Supabase CLI.
"""

from __future__ import annotations

import shutil
import subprocess
import tempfile
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


PROJECT_REF = "pnhskmlejdaklwkvasxp"  # BASE GCBA
SCHEMA = "red_conexion_gerencial"
SHEET = "Respuestas de formulario 1"
ROOT = Path(__file__).resolve().parents[1]
COLUMNS = (
    "source_sheet",
    "source_row",
    "source_registered_at",
    "given_name",
    "family_name",
    "position_title",
    "ministry",
    "secretariat",
    "directorate",
    "phone",
    "email",
    "advisory_topics",
)


def sql_value(value: object) -> str:
    if value is None:
        return "NULL"
    if isinstance(value, datetime):
        value = value.isoformat(sep=" ")
    return "'" + str(value).replace("\x00", "").replace("'", "''") + "'"


def trimmed(value: object) -> str | None:
    if value is None:
        return None
    result = str(value).strip()
    return result or None


def main() -> None:
    sources = list(ROOT.glob("*.xlsx"))
    if len(sources) != 1:
        raise SystemExit("Se esperaba exactamente un archivo XLSX en la raíz.")

    workbook = load_workbook(sources[0], read_only=True, data_only=True)
    if SHEET not in workbook.sheetnames:
        raise SystemExit("No se encontró la hoja de respuestas del directorio.")
    sheet = workbook[SHEET]
    headers = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
    if trimmed(headers[1]) != "Apellido" or trimmed(headers[3]) != "Nombre":
        raise SystemExit("Las columnas del XLSX no coinciden con el importador.")
    if "Email" not in str(headers[9]) or "Cuit" not in str(headers[11]):
        raise SystemExit("La estructura del XLSX cambió; se canceló la importación.")

    records: list[tuple[object, ...]] = []
    seen_emails: set[str] = set()
    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
        if not any(value is not None and str(value).strip() for value in row):
            continue
        email = trimmed(row[9])
        if not email or email.count("@") != 1 or email.lower() in seen_emails:
            raise SystemExit(f"Email faltante, inválido o repetido en la fila {row_number}.")
        seen_emails.add(email.lower())
        required = (trimmed(row[3]), trimmed(row[1]), trimmed(row[4]), trimmed(row[5]))
        if any(value is None for value in required):
            raise SystemExit(f"Falta un campo obligatorio en la fila {row_number}.")
        records.append(
            (
                SHEET,
                row_number,
                row[0] if isinstance(row[0], datetime) else None,
                required[0],
                required[1],
                required[2],
                required[3],
                trimmed(row[6]),
                trimmed(row[7]),
                trimmed(row[8]) or "",
                email,
                trimmed(row[10]),
            )
        )
    workbook.close()

    cli = shutil.which("npx")
    if not cli:
        raise SystemExit("No se encontró npx para ejecutar Supabase CLI.")
    print(f"Registros válidos preparados: {len(records)}. Destino: BASE GCBA/{SCHEMA}.people")

    for offset in range(0, len(records), 75):
        batch = records[offset : offset + 75]
        values = ",\n".join("(" + ", ".join(sql_value(value) for value in record) + ")" for record in batch)
        statement = (
            "begin;\n"
            f"insert into {SCHEMA}.people ({', '.join(COLUMNS)}) values\n{values}\n"
            "on conflict do nothing;\n"
            "commit;\n"
        )
        temporary_path: Path | None = None
        try:
            with tempfile.NamedTemporaryFile(
                mode="w", encoding="utf-8", suffix=".sql", prefix="rcg-import-", delete=False
            ) as temporary:
                temporary.write(statement)
                temporary_path = Path(temporary.name)
            completed = subprocess.run(
                [
                    cli,
                    "--yes",
                    "supabase@latest",
                    "db",
                    "query",
                    "--linked",
                    "--project-ref",
                    PROJECT_REF,
                    "--file",
                    str(temporary_path),
                    "--agent",
                    "no",
                    "--output-format",
                    "text",
                ],
                cwd=ROOT,
                capture_output=True,
                text=True,
                check=False,
            )
            if completed.returncode:
                raise SystemExit(
                    f"Falló el lote {offset // 75 + 1} (código {completed.returncode}). "
                    "No se muestran mensajes SQL para proteger los datos personales."
                )
        finally:
            if temporary_path is not None:
                temporary_path.unlink(missing_ok=True)
        print(f"Lote {offset // 75 + 1}: procesados {len(batch)} registros.")


if __name__ == "__main__":
    main()
