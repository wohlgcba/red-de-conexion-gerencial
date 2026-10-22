"""Offline authority reconciliation. Never writes to Supabase or the source roster.

Only private staging under .local may contain linking identifiers. Source periods
are explicit inputs, never inferred from file modification times or seniority.
"""
from __future__ import annotations
import argparse
import csv
from datetime import date, datetime
import hashlib
import json
from pathlib import Path
import re
from collections import Counter, defaultdict
from openpyxl import load_workbook
from build_organization_sample import norm, valid_key

ROOT = Path(__file__).resolve().parents[1]
LOCAL = ROOT / '.local'
SNAPSHOT_SQL = """select jsonb_build_object(
 'units',(select jsonb_agg(to_jsonb(u)) from red_conexion_gerencial.organization_units u),
 'assignments',(select jsonb_agg(to_jsonb(a)||jsonb_build_object('link_key',o.link_key,'name',o.institutional_name))
  from rcg_private.organization_assignments a join rcg_private.organization_people o on o.id=a.responsible_id)
) as snapshot;"""
TARGET_TYPES = {'organismo', 'secretaria', 'subsecretaria', 'direccion general',
                'gerencia general', 'gerencia operativa', 'subgerencia operativa'}


def normalized(value):
    return re.sub(r'[^a-z0-9]+', ' ', norm(value)).strip()


def canonical_name(value):
    text = normalized(value)
    for short, full in [('ss', 'subsecretaria'), ('sec', 'secretaria'), ('secr', 'secretaria'),
                        ('dg', 'direccion general'), ('go', 'gerencia operativa'),
                        ('sgo', 'subgerencia operativa'), ('gg', 'gerencia general')]:
        text = re.sub('^' + short + r'\b', full, text)
    return re.sub(r'\s+', ' ', re.sub(r'\b(?:de|del|la|el|los|las)\b', '', text)).strip()


def verified_names(unit):
    return {canonical_name(unit['name'])} | {canonical_name(h['name']) for h in unit.get('name_history', []) if isinstance(h, dict) and h.get('name')}


def code_key(value):
    """Compare numeric magnitude without inventing a width or altering the code."""
    text = str(value or '').strip()
    if re.fullmatch(r'\d+\.0', text):
        text = text[:-2]
    if re.fullmatch(r'\d+', text):
        return ('numeric', str(int(text))) if int(text) else None
    return ('literal', text.upper()) if text else None


def parse_date(value):
    if value in (None, ''):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    for fmt in ('%d/%m/%Y', '%Y-%m-%d', '%Y-%m-%d %H:%M:%S'):
        try:
            return datetime.strptime(str(value).strip(), fmt).date()
        except ValueError:
            pass
    raise ValueError('Fecha institucional inválida; requiere revisión de la fuente.')


def role_kind(role):
    """Exact institutional enums observed in Dotación, not generic word matches."""
    role = normalized(role)
    return {
        'jefe de gobierno': 'government', 'vicejefe de gobierno': 'vicegovernment',
        'jefe de gabinete de ministros': 'cabinet', 'ministro': 'minister',
        'secretario': 'secretary', 'secretario especial': 'secretary',
        'secretario legal y tecnico': 'secretary', 'subsecretario': 'undersecretary',
        'director general': 'director', 'gerente general sindicatura': 'general_manager',
        'gerente general': 'general_manager', 'gerente operativo': 'manager',
        'gerente operativo concursado': 'manager', 'subgerente operativo': 'submanager',
        'titular de unidad': 'unit_head', 'titular del organismo': 'agency_head',
        'presidente cpc': 'president', 'director ejecutivo': 'executive',
        'sindico general': 'syndic', 'sindico general adjunto': 'deputy_syndic',
        'procurador general': 'procurator', 'procurador general adjunto': 'deputy_procurator',
    }.get(role)


def compatible(unit, row):
    kind = role_kind(row['role'])
    name, typ = canonical_name(unit['name']), normalized(unit['unit_type'])
    # Administrative rank differs from the actual institutional office here.
    if name.startswith('sindicatura general'):
        if 'adjunt' in name and kind in ('undersecretary', 'deputy_syndic'):
            return 'Síndico/a General Adjunto/a'
        if 'adjunt' not in name and kind in ('minister', 'syndic'):
            return 'Síndico/a General'
        return None
    if name.startswith('procuracion general') or name.startswith('pg adjunta'):
        if 'adjunt' in name and kind in ('undersecretary', 'deputy_procurator'):
            return 'Procurador/a General Adjunto/a'
        if 'adjunt' not in name and kind in ('minister', 'procurator'):
            return 'Procurador/a General'
        return None
    if name == 'jefe gobierno' and kind == 'government':
        return 'Jefe de Gobierno'
    if name == 'vicejefatura gobierno' and kind == 'vicegovernment':
        return 'Vicejefe/a de Gobierno'
    if name == 'jefatura gabinete ministros' and kind == 'cabinet':
        return 'Jefe de Gabinete de Ministros'
    if name.startswith('ministerio ') and kind == 'minister':
        return 'Ministro/a'
    expected = {'subgerencia operativa': 'submanager', 'gerencia operativa': 'manager',
                'gerencia general': 'general_manager', 'direccion general': 'director',
                'subsecretaria': 'undersecretary', 'secretaria': 'secretary'}
    if expected.get(typ) == kind and kind:
        return {'submanager': 'Subgerente Operativo/a', 'manager': 'Gerente Operativo/a',
                'general_manager': 'Gerente General', 'director': 'Director/a General',
                'undersecretary': 'Subsecretario/a', 'secretary': 'Secretario/a'}[kind]
    if kind == 'unit_head' and name.startswith('unidad '):
        return 'Titular de la unidad'
    if kind in ('agency_head', 'president', 'executive') and typ in ('organismo', 'unidad organizacional'):
        return {'agency_head': 'Titular del organismo', 'president': 'Presidente/a', 'executive': 'Director/a Ejecutivo/a'}[kind]
    return None


def read_source(path, period, official=False, reference=None):
    """Whitelist identity and institutional fields; discard all contacts/demographics."""
    path = Path(path)
    if official:
        required = {'ayn', 'cuil', 'lit_puesto', 'cod_rep', 'desc_rep', 'ministerio', 'reporta_a', 'car_sit_rev', 'rol_desde', 'rol_hasta'}
    else:
        required = {'AYN', 'CUIL', 'LIT_PUESTO', 'COD_REP', 'DESC_REP', 'MINISTERIO', 'PATH_NOMBRES'}
    book = None
    if path.suffix.lower() == '.csv':
        text = path.read_text(encoding='utf-8-sig')
        dialect = csv.Sniffer().sniff(text[:4000], delimiters=',;\t')
        reader = csv.DictReader(text.splitlines(), dialect=dialect)
        columns, records = set(reader.fieldnames or []), reader
    else:
        book = load_workbook(path, read_only=True, data_only=True)
        iterator = book.worksheets[0].iter_rows(values_only=True)
        headers = [str(value or '').strip() for value in next(iterator)]
        columns, records = set(headers), (dict(zip(headers, row)) for row in iterator)
    if not required.issubset(columns):
        if book: book.close()
        raise ValueError('La fuente no contiene el esquema institucional requerido.')
    result, units, roles = [], {}, Counter()
    try:
        for record in records:
            get = lambda key: record.get(key.lower() if official else key)
            raw_code, ministry, unit = get('COD_REP'), str(get('MINISTERIO') or '').strip(), str(get('DESC_REP') or '').strip()
            path_text = str(record.get('reporta_a' if official else 'PATH_NOMBRES') or '')
            code = str(raw_code or '').strip()
            units[(code, ministry, path_text)] = unit
            role = str(get('LIT_PUESTO') or '').strip()
            roles[role] += 1
            if not role_kind(role):
                continue
            name = str(get('AYN') or '').strip()
            if ',' in name:
                family, given = name.split(',', 1)
                name = (given.strip() + ' ' + family.strip()).title()
            since, until, invalid_date = None, None, False
            try:
                if official:
                    since, until = parse_date(record.get('rol_desde')), parse_date(record.get('rol_hasta'))
            except ValueError:
                invalid_date = True
            result.append({'key': valid_key(get('CUIL')), 'name': name, 'role': role,
                'code': code, 'unit': unit, 'ministry': ministry, 'path': path_text,
                'source': 'Autoridades oficiales 2026' if official else 'Dotación institucional',
                'source_period': period, 'source_date': reference,
                'since': since.isoformat() if since else None, 'until': until.isoformat() if until else None,
                'active': normalized(record.get('car_sit_rev')) if official else None,
                'invalid_date': invalid_date, 'official': official,
                'category': str(record.get('escalafon' if official else 'REGIMEN') or ''),
                'sigla': str(get('SIGLA') or '')})
    finally:
        if book: book.close()
    return result, [{'code': k[0], 'ministry': k[1], 'path': k[2], 'name': v} for k, v in units.items()], roles


def temporal(row, as_of):
    if row.get('source_date') and date.fromisoformat(row['source_date']) > as_of:
        return False, 'Fuente posterior a la fecha de evaluación'
    if row.get('invalid_date'):
        return False, 'Fechas de vigencia inválidas'
    if row.get('since') and date.fromisoformat(row['since']) > as_of:
        return False, 'Designación futura'
    if row.get('until') and date.fromisoformat(row['until']) < as_of:
        return False, 'Designación finalizada'
    if row.get('official'):
        if row.get('active') != 'activo':
            return False, 'Registro histórico o situación no confirmada como activa'
        if not row.get('since') or not row.get('until'):
            return None, 'Falta confirmar vigencia de la designación'
        if row.get('source_period') != as_of.strftime('%Y-%m'):
            return None, 'Fuente oficial de otro período; requiere actualización'
        return True, 'Intervalo de vigencia y situación activa verificados'
    # A current institutional roster is evidence of presence in that MONTH,
    # not an invented start/end date, nor an assertion of permanent tenure.
    if row.get('source_period') == as_of.strftime('%Y-%m'):
        return True, 'Presencia en Dotación del mes de referencia; sin fechas individuales disponibles'
    return None, 'Dotación de otro período; requiere verificación de vigencia'


def reconcile(snapshot, records, as_of):
    units = {u['id']: u for u in snapshot['units']}
    paths = {}
    def path_for(unit):
        if unit['id'] in paths: return paths[unit['id']]
        parent = units.get(unit['parent_id'])
        paths[unit['id']] = (path_for(parent) if parent else []) + [unit]
        return paths[unit['id']]
    by_code, by_name = defaultdict(list), defaultdict(list)
    for row in records:
        if code_key(row['code']): by_code[code_key(row['code'])].append(row)
        by_name[canonical_name(row['unit'])].append(row)
    existing = defaultdict(list)
    for item in snapshot.get('assignments') or []: existing[item['unit_id']].append(item)
    results = []
    for unit in units.values():
        path = path_for(unit)
        candidates, seen = [], set()
        # A homonymous office elsewhere must not compete with an exact code.
        pool = by_code.get(code_key(unit.get('code')), []) or [row for name in verified_names(unit) for row in by_name.get(name, [])]
        for row in pool:
            identity = (row.get('key'), row['role'], row['source'], row['code'], row['path'], row.get('since'), row.get('until'))
            if identity in seen: continue
            seen.add(identity)
            role = compatible(unit, row)
            if not role: continue
            exact_code = bool(code_key(unit.get('code'))) and code_key(unit['code']) == code_key(row['code'])
            root_ok = canonical_name(row['ministry']) in verified_names(path[0])
            names = [canonical_name(x) for x in row['path'].split('\\') if x.strip()]
            context_ok = len(names) == len(path) and all(name in verified_names(ancestor) for name, ancestor in zip(names, path))
            leaf_ok = canonical_name(row['unit']) in verified_names(unit)
            current, validity = temporal(row, as_of)
            name_valid = bool(row['name'].strip()) and not any(char in row['name'] for char in ('?', '\ufffd'))
            confidence = 'high' if exact_code and root_ok and context_ok and leaf_ok and current is True and row.get('key') and name_valid else 'medium'
            if not root_ok or current is False: confidence = 'low'
            reasons = []
            if not exact_code: reasons.append('Falta código institucional coincidente')
            if not root_ok: reasons.append('Organismo incompatible')
            if not context_ok: reasons.append('Trayectoria jerárquica sin confirmar')
            if not leaf_ok: reasons.append('Denominación distinta: posible cambio de nombre o de dependencia')
            if not row.get('key'): reasons.append('Identidad no validada')
            if not name_valid: reasons.append('Nombre incompleto o con caracteres dañados en la fuente')
            if current is not True: reasons.append(validity)
            candidates.append({**row, 'institutional_role': role, 'confidence': confidence,
                'method': 'code_exact' if str(unit.get('code')) == row['code'] else 'code_numeric_equivalent' if exact_code else 'contextual_name',
                'reason': '; '.join(reasons) or validity, 'eligible': root_ok and current is not False and row.get('source_period') == as_of.strftime('%Y-%m')})
        # A repeated row/person is not an ambiguity. Distinct active identities are.
        eligible = [c for c in candidates if c['eligible'] and c.get('key')]
        identities = {c['key'] for c in eligible}
        high = [c for c in eligible if c['confidence'] == 'high']
        # A monthly roster cannot silently overrule explicit end dates or an
        # inactive designation in the complementary source from the same period.
        official_active = {c['key'] for c in candidates if c.get('official') and temporal(c, as_of)[0] is True}
        temporal_conflict = any(c.get('official') and c.get('key') in {h['key'] for h in high}
            and c.get('key') not in official_active and c.get('source_period') == as_of.strftime('%Y-%m')
            and temporal(c, as_of)[0] is False for c in candidates)
        old = existing.get(unit['id'], [])
        manual = [a for a in old if a['responsibility_confirmed'] and
                  (a.get('assignment_origin') == 'manual' or
                   a.get('assignment_origin') not in ('automatic',) and not str(a.get('verification_basis') or '').startswith('Único candidato con identificador válido y cargo jerárquico compatible'))]
        status, winner, reason = 'unidentified', None, 'Sin evidencia suficiente de titularidad'
        if manual or unit.get('authority_reviewed'):
            status, reason = unit.get('authority_status', 'manual'), 'Decisión manual previa preservada; no se sobrescribe automáticamente'
            if manual: status = 'manual'
            if manual and (not identities or any(a.get('link_key') not in identities for a in manual)): reason += '; discrepancia con las fuentes actuales'
        elif len(identities) > 1:
            status, reason = 'pending', 'Varios candidatos activos o fuentes en conflicto'
        elif any(c['eligible'] and not c.get('key') for c in candidates):
            status, reason = 'pending', 'Hay candidatos cuya identidad no pudo validarse'
        elif temporal_conflict:
            status, reason = 'pending', 'La vigencia de la fuente complementaria contradice la presencia en Dotación; requiere revisión'
        elif high and len(identities) == 1:
            winner = sorted(high, key=lambda c: c['official'])[0]
            status, reason = 'automatic', winner['reason']
        elif candidates:
            status = 'pending'
            reason = '; '.join(dict.fromkeys(c['reason'] for c in candidates))
        if normalized(unit['name']) == 'obligaciones a cargo del tesoro' and not winner:
            reason = 'Posible clasificación presupuestaria: no se presume un titular propio sin evidencia institucional'
        results.append({'unit_id': unit['id'], 'name': unit['name'], 'type': unit['unit_type'],
            'organism': path[0]['name'], 'code': unit.get('code'), 'status': status,
            'winner': winner, 'candidates': candidates, 'reason': reason,
            'expected': normalized(unit['unit_type']) in TARGET_TYPES or bool(candidates or manual),
            'before': any(a['responsibility_confirmed'] for a in old)})
    return results


def coverage(results):
    def counts(items):
        # Scope is NOT proof that a budget classification has its own holder.
        return {'total': len(items), 'in_scope': sum(r['expected'] for r in items),
                'holder_requirement_confirmed': sum(r['status'] in ('automatic','manual','vacant') for r in items),
                'holder_requirement_unknown': sum(r['expected'] and r['status'] in ('pending','unidentified') for r in items),
                'before': sum(r['before'] for r in items),
                **{state: sum(r['status'] == state for r in items) for state in ('automatic','manual','pending','vacant','unidentified','not_applicable')}}
    return {'total': counts(results), 'by_type': {name: counts([r for r in results if r['type'] == name]) for name in sorted({r['type'] for r in results})},
        'by_organism': {name: counts([r for r in results if r['organism'] == name]) for name in sorted({r['organism'] for r in results})}}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--dotacion', type=Path, required=True)
    parser.add_argument('--period', required=True, help='Verified institutional roster period YYYY-MM')
    parser.add_argument('--as-of', type=date.fromisoformat, required=True)
    parser.add_argument('--official', type=Path)
    parser.add_argument('--official-date', type=date.fromisoformat)
    parser.add_argument('--offline', action='store_true', help='Use the already-read private Supabase snapshot')
    parser.add_argument('--renames', type=Path, help='Explicit evidence-backed rename proposals; never modifies the hierarchy')
    args = parser.parse_args()
    if not re.fullmatch(r'\d{4}-\d{2}', args.period): raise ValueError('Período inválido')
    LOCAL.mkdir(exist_ok=True)
    target = LOCAL / 'authority-current-private.json'
    if not args.offline:
        from supabase_project import read_query
        target.write_text(json.dumps(read_query(SNAPSHOT_SQL)[0]['snapshot']), encoding='utf8')
    snapshot = json.loads(target.read_text(encoding='utf8'))
    records, source_units, roles = read_source(args.dotacion, args.period)
    if args.official:
        if args.official_date is None: raise ValueError('La fuente oficial requiere fecha de referencia explícita')
        if args.official_date.strftime('%Y-%m') != args.period: raise ValueError('No se pueden mezclar períodos de las fuentes institucionales')
        extra, _, _ = read_source(args.official, args.official_date.strftime('%Y-%m'), True, args.official_date.isoformat())
        records.extend(extra)
    results = reconcile(snapshot, records, args.as_of)
    if args.renames:
        apply_rename_proposals(results, snapshot, json.loads(args.renames.read_text(encoding='utf8')), args.as_of)
    private = {'as_of': args.as_of.isoformat(), 'period': args.period, 'results': results,
               'records': records, 'source_units': source_units, 'roles': dict(roles), 'source_sha256': hashlib.sha256(args.dotacion.read_bytes()).hexdigest()}
    (LOCAL / 'authority-reconciliation-private.json').write_text(json.dumps(private, ensure_ascii=False), encoding='utf8')
    report = coverage(results)
    (LOCAL / 'authority-coverage.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf8')
    # Compact unresolved report intentionally excludes people and identifiers.
    with (LOCAL / 'authority-unresolved.csv').open('w', encoding='utf-8-sig', newline='') as file:
        writer = csv.writer(file); writer.writerow(['Dependencia', 'Tipo', 'Organismo', 'Estado', 'Motivo'])
        writer.writerows([r['name'], r['type'], r['organism'], r['status'], r['reason']] for r in results if r['expected'] and r['status'] not in ('automatic', 'manual'))
    # Source-only offices are evidence for structural review, never new nodes.
    known_codes = {code_key(u.get('code')) for u in snapshot['units']} - {None}
    missing = {(r['code'], r['unit'], r['ministry']) for r in records if code_key(r['code']) and code_key(r['code']) not in known_codes and temporal(r,args.as_of)[0] is not False}
    with (LOCAL / 'authority-source-only-units.csv').open('w', encoding='utf-8-sig', newline='') as file:
        writer = csv.writer(file); writer.writerow(['Código fuente','Dependencia fuente','Organismo','Motivo'])
        writer.writerows([*unit, 'Código ausente en la estructura actual; revisar alta, cambio o código sin confirmar'] for unit in sorted(missing))
    print(json.dumps(report, ensure_ascii=True))


def apply_rename_proposals(results, snapshot, proposals, as_of):
    """A different label is a review proposal, never proof of institutional continuity."""
    units = {u['id']: u for u in snapshot['units']}
    indexed = {r['unit_id']: r for r in results}
    for proposal in proposals:
        unit = units.get(proposal['unit_id'])
        if (not unit or unit.get('code') != proposal.get('code') or unit['name'] != proposal.get('previous_name')
            or indexed[unit['id']]['organism'] != proposal.get('organism')
            or not proposal.get('new_name') or not proposal.get('evidence')
            or date.fromisoformat(proposal['observed_at']) > as_of):
            raise ValueError('Propuesta de denominación sin contexto institucional y evidencia suficientes')
        indexed[unit['id']]['rename'] = {'name': proposal['new_name'], 'evidence': proposal['evidence'], 'observed_at': proposal['observed_at']}


if __name__ == '__main__': main()
