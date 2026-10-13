"""Read-only Dotación analysis and offline reconciliation; no remote operations.

Only institutional units and aggregate counts can be exported into src/mocks.
Personal identifiers are processed in memory, never logged or written to src.
Usage: python scripts/build_organization_sample.py DOTACION --directory DIRECTORIO
Requires openpyxl (already used by the directory importer).
"""
from __future__ import annotations
import argparse
import json
import re
import sqlite3
import unicodedata
import uuid
from collections import Counter, defaultdict
from pathlib import Path
from openpyxl import load_workbook

NAMESPACE = uuid.UUID('73aa6756-3ec0-44ab-8869-bc26930227c8')
ROOT = Path(__file__).resolve().parents[1]

def norm(value):
    return ''.join(c for c in unicodedata.normalize('NFKD', str(value or '').casefold()) if not unicodedata.combining(c))

def valid_key(value):
    if isinstance(value, float):
        if not value.is_integer(): return None
        value = int(value)
    text = re.sub(r'[\s-]', '', str(value or ''))
    if not re.fullmatch(r'\d{11}', text): return None
    remainder = 11 - sum(int(a) * b for a, b in zip(text[:10], [5,4,3,2,7,6,5,4,3,2])) % 11
    check = 0 if remainder == 11 else 9 if remainder == 10 else remainder
    return text if int(text[-1]) == check else None

def unit_type(name, root):
    text = norm(name)
    if root: return 'Organismo'
    for prefixes, label in [(['subgerencia', 'sgo '], 'Subgerencia Operativa'),
        (['gerencia general'], 'Gerencia General'), (['gerencia operativa', 'go '], 'Gerencia Operativa'),
        (['subdireccion general', 'sdg '], 'Subdirección General'),
        (['direccion general', 'dg '], 'Dirección General'), (['subsecretar', 'ss '], 'Subsecretaría'),
        (['secretar', 'secr '], 'Secretaría'), (['direccion', 'dir ', 'dir. '], 'Dirección'),
        (['subdireccion'], 'Subdirección'), (['dept ', 'departamento '], 'Departamento'),
        (['div ', 'division '], 'División'), (['secc ', 'seccion '], 'Sección')]:
        if any(text.startswith(prefix) for prefix in prefixes): return label
    return 'Unidad organizacional'

def directory_keys(path):
    keys = Counter()
    records = []
    stats = {'directoryRows': 0, 'directoryInvalidKeys': 0}
    if not path: return keys, stats, records
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = book['Respuestas de formulario 1'].iter_rows(values_only=True)
        headers = next(rows)
        columns = [i for i, header in enumerate(headers) if 'cuit' in norm(header) or 'cuil' in norm(header)]
        emails = [i for i, header in enumerate(headers) if 'email' in norm(header)]
        if len(columns) != 1: raise ValueError('No se identificó una columna única de conciliación en el Directorio.')
        for source_row, row in enumerate(rows, 2):
            if not any(row): continue
            stats['directoryRows'] += 1
            key = valid_key(row[columns[0]])
            if key:
                keys[key] += 1
                records.append((key, source_row, str(row[emails[0]] or '').strip().casefold() if len(emails) == 1 else ''))
            else: stats['directoryInvalidKeys'] += 1
    finally: book.close()
    stats['directoryDuplicateKeys'] = sum(n > 1 for n in keys.values())
    return keys, stats, records

def write_private_review(path, units, directory_records, candidate_records):
    # A review database may contain identifiers; keep it outside public assets
    # and inside the ignored local workspace directory. It is NEVER an API DTO.
    path = path.resolve()
    local_root = (ROOT / '.local').resolve()
    if not path.is_relative_to(local_root): raise ValueError('La revisión privada debe guardarse dentro de .local.')
    path.parent.mkdir(parents=True, exist_ok=True)
    with sqlite3.connect(path) as db:
        db.executescript('''
            create table if not exists units(id text primary key, parent_id text, name text, unit_type text, code text);
            create table if not exists directory_keys(link_key text, source_row integer, expected_email text, primary key(link_key, source_row));
            create table if not exists candidates(unit_id text, link_key text, institutional_name text, position_title text,
                review_state text not null default 'pending', primary key(unit_id,link_key));
        ''')
        # Preserve reviewed decisions. Update only the minimal source snapshot.
        db.executemany('insert or replace into units values(?,?,?,?,?)', [(u['id'],u['parentId'],u['name'],u['type'],u['code']) for u in units])
        db.executemany('insert or replace into directory_keys values(?,?,?)', directory_records)
        db.executemany('insert into candidates(unit_id,link_key,institutional_name,position_title) values(?,?,?,?) on conflict(unit_id,link_key) do update set institutional_name=excluded.institutional_name,position_title=excluded.position_title', candidate_records)
        matches = db.execute('''select count(distinct c.link_key) from candidates c
            join (select link_key from directory_keys group by link_key having count(*)=1) d using(link_key)''').fetchone()[0]
    return matches

def extract(path, directory=None, sheet_name=None, review_db=None):
    directory_index, directory_stats, directory_records = directory_keys(directory)
    book = load_workbook(path, read_only=True, data_only=True)
    units, codes, cache = {}, defaultdict(set), {}
    candidates = defaultdict(set)
    candidate_records = {}
    valid_people, matches = set(), set()
    invalid = rows_count = mismatches = no_path = 0
    try:
        sheet = book[sheet_name] if sheet_name else book.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        headers = [str(value or '').strip().upper() for value in next(rows)]
        required = {'PATH_NOMBRES', 'MINISTERIO', 'DESC_REP', 'COD_REP', 'SIGLA', 'LIT_PUESTO', 'CUIL', 'AYN'}
        if not required.issubset(headers): raise ValueError('Faltan columnas requeridas de Dotación.')
        columns = {name: headers.index(name) for name in required}
        for row in rows:
            if not any(row): continue
            rows_count += 1
            raw_path = str(row[columns['PATH_NOMBRES']] or '').strip()
            if not raw_path: no_path += 1; continue
            if raw_path not in cache:
                names = [re.sub(r'\s+', ' ', name).strip() for name in raw_path.split('\\') if name.strip()]
                parent = None
                for depth, name in enumerate(names):
                    context = '\\'.join(names[:depth + 1]).casefold()
                    identifier = str(uuid.uuid5(NAMESPACE, context))
                    units.setdefault(identifier, {'id': identifier, 'parentId': parent, 'name': name,
                        'type': unit_type(name, depth == 0), 'code': None, 'directorateId': None})
                    parent = identifier
                cache[raw_path] = (parent, names[0] if names else '')
            identifier, root_name = cache[raw_path]
            if not identifier: no_path += 1; continue
            if norm(row[columns['MINISTERIO']]).strip() != norm(root_name).strip(): mismatches += 1
            code = row[columns['COD_REP']]
            if code is not None: codes[identifier].add(str(code).strip())
            key = valid_key(row[columns['CUIL']])
            if key:
                valid_people.add(key)
                if directory_index.get(key) == 1: matches.add(key)
            else: invalid += 1
            role = norm(row[columns['LIT_PUESTO']]).strip()
            # Candidates need later confirmation: belonging to an area is not proof.
            expected = {'Subgerencia Operativa': ('subgerente operativo',), 'Gerencia Operativa': ('gerente operativo',),
                'Gerencia General': ('gerente general',), 'Dirección General': ('director general', 'directora general'),
                'Secretaría': ('secretario', 'secretaria'), 'Subsecretaría': ('subsecretario', 'subsecretaria')}.get(units[identifier]['type'], ())
            if key and str(row[columns['AYN']] or '').strip() and any(role == prefix or role.startswith(prefix + ' ') for prefix in expected):
                candidates[identifier].add(key)
                if review_db:
                    candidate_records[(identifier, key)] = (identifier, key, str(row[columns['AYN']]).strip(), str(row[columns['LIT_PUESTO']]).strip())
    finally: book.close()
    for identifier, values in codes.items():
        if len(values) == 1: units[identifier]['code'] = next(iter(values))
    stats = {'filename': path.name, 'sheet': sheet.title, 'rows': rows_count, 'units': len(units),
        'roots': sum(unit['parentId'] is None for unit in units.values()), 'rootDisagreements': mismatches,
        'rowsWithoutPath': no_path, 'codeConflicts': sum(len(values) > 1 for values in codes.values()),
        'validDistinctPeople': len(valid_people), 'invalidIdentifierRows': invalid,
        'directoryMatchesOffline': len(matches), 'candidateUnits': len(candidates),
        'ambiguousCandidateUnits': sum(len(values) > 1 for values in candidates.values()),
        'candidatePeople': len(set().union(*candidates.values())) if candidates else 0,
        'responsibilities': 'not published; actual candidates require server verification', **directory_stats}
    if review_db:
        stats['candidateDirectoryMatchesOffline'] = write_private_review(review_db, list(units.values()), directory_records, list(candidate_records.values()))
    return {'units': list(units.values()), 'sample': stats}

if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('xlsx', type=Path)
    parser.add_argument('--directory', type=Path)
    parser.add_argument('--sheet')
    parser.add_argument('--review-db', type=Path, help='Private SQLite review file inside .local; never imported remotely.')
    args = parser.parse_args()
    result = extract(args.xlsx, args.directory, args.sheet, args.review_db)
    target = ROOT / 'src/mocks/organizationStructure.json'
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(json.dumps({**result['sample'], 'personalFieldsExported': False}, ensure_ascii=True))
